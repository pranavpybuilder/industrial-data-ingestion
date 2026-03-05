"""
DOCX Exporter — MODE 1 Insights Report (.docx)

Generates a professional Word document with:
- Cover page with title and date
- Auto-generated Table of Contents
- Executive Summary styled box
- Data Overview table
- Key Findings with severity colors
- Root Cause Analysis table
- 3-Tier Recommendations (red/amber/blue headers)
- Predictive Signals section
- Footer: "Confidential | Offline Industrial Intelligence | {date}"

Dependencies: python-docx
"""

import os
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from utils.paths import EXPORT_DIR
from storage.connection import get_run_file_name
from utils.logger import get_logger

logger = get_logger(__name__)

# ──────────────────────────────────────────
# Color Constants
# ──────────────────────────────────────────
NAVY = RGBColor(0x1F, 0x4E, 0x78)
DARK_BLUE = RGBColor(0x36, 0x60, 0x92)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x73, 0x06)
BLUE = RGBColor(0x25, 0x63, 0xEB)
GREEN = RGBColor(0x05, 0x96, 0x69)
GRAY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _set_cell_bg(cell, hex_color: str) -> None:
    """
    Safely set a table cell background color.
    Compatible with python-docx >= 1.2.0 (no get_or_add_tcPr).

    This function manually finds or creates the w:tcPr element
    and appends a w:shd child, avoiding the removed get_or_add API.
    """
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color.replace('#', ''))
    tcPr.append(shd)


class DocxExporter:
    """
    Exports insights and profiling results to a professional Word document.
    """

    def __init__(self) -> None:
        self.logger = logger

    def export_insights_docx(
        self,
        run_id: str,
        unified_insights: List[Dict[str, Any]],
        profiling_results: Optional[Dict[str, Any]] = None,
        narrative: Optional[Dict[str, Any]] = None,
        ml_findings: Optional[List[Dict[str, Any]]] = None,
        output_dir: Optional[str] = None,
    ) -> str:
        """
        Generate a complete DOCX insights report.
        Returns the absolute file path of the generated .docx file.
        """
        target_dir = Path(output_dir) if output_dir else EXPORT_DIR / run_id
        target_dir.mkdir(parents=True, exist_ok=True)

        file_stem = get_run_file_name(run_id)
        date_tag = datetime.now().strftime("%Y%m%d")
        filename = str(target_dir / f"{file_stem}_insights_report_{date_tag}.docx")

        doc = Document()
        self._apply_default_font(doc)

        self._add_cover_page(doc, run_id)
        self._add_toc(doc)
        doc.add_page_break()
        self._add_executive_summary(doc, unified_insights, narrative)
        self._add_data_overview(doc, unified_insights, profiling_results)
        self._add_key_findings(doc, unified_insights)
        self._add_root_cause_table(doc, unified_insights)
        self._add_recommendations(doc, unified_insights)
        self._add_predictive_signals(doc, ml_findings or [])
        self._add_footer(doc)

        doc.save(filename)
        self.logger.info(f"DOCX insights report exported: {filename}")
        return filename

    # ────────────────────────────────────────
    # Private Section Builders
    # ────────────────────────────────────────

    @staticmethod
    def _apply_default_font(doc: Document) -> None:
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    def _add_cover_page(self, doc: Document, run_id: str) -> None:
        for _ in range(6):
            doc.add_paragraph("")

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("Industrial Intelligence Report")
        run.font.size = Pt(32)
        run.font.color.rgb = NAVY
        run.bold = True

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run("Offline Analysis & Insights")
        run.font.size = Pt(18)
        run.font.color.rgb = DARK_BLUE

        doc.add_paragraph("")

        details = doc.add_paragraph()
        details.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = details.add_run(f"Run ID: {run_id}")
        run.font.size = Pt(12)
        run.font.color.rgb = GRAY

        date_line = doc.add_paragraph()
        date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = date_line.add_run(datetime.now().strftime("%B %d, %Y"))
        run.font.size = Pt(12)
        run.font.color.rgb = GRAY

        classification = doc.add_paragraph()
        classification.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = classification.add_run("CONFIDENTIAL")
        run.font.size = Pt(10)
        run.font.color.rgb = RED
        run.bold = True

        doc.add_page_break()

    @staticmethod
    def _add_toc(doc: Document) -> None:
        heading = doc.add_heading("Table of Contents", level=1)
        heading.runs[0].font.color.rgb = NAVY

        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        fld_char_begin = run._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
        run._element.append(fld_char_begin)

        run2 = paragraph.add_run()
        instr_text = run2._element.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"})
        instr_text.text = ' TOC \\o "1-3" \\h \\z \\u '
        run2._element.append(instr_text)

        run3 = paragraph.add_run()
        fld_char_end = run3._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
        run3._element.append(fld_char_end)

        note = doc.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = note.add_run("(Right-click \u2192 Update Field to refresh TOC)")
        run.font.size = Pt(9)
        run.font.color.rgb = GRAY
        run.italic = True

    def _add_executive_summary(
        self,
        doc: Document,
        insights: List[Dict[str, Any]],
        narrative: Optional[Dict[str, Any]],
    ) -> None:
        heading = doc.add_heading("Executive Summary", level=1)
        heading.runs[0].font.color.rgb = NAVY

        critical = sum(1 for i in insights if str(i.get("severity", "")).upper() == "CRITICAL")
        warnings = sum(1 for i in insights if str(i.get("severity", "")).upper() == "WARNING")
        info_count = len(insights) - critical - warnings

        # Summary box using safe cell background
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)

        # SAFE: use _set_cell_bg instead of broken get_or_add API
        _set_cell_bg(cell, 'EFF6FF')

        if narrative and isinstance(narrative, dict):
            sections = narrative.get("sections", narrative)
            exec_summary = sections.get("Executive Summary", "")
            if exec_summary:
                cell.text = str(exec_summary)
            else:
                cell.text = self._build_fallback_summary(len(insights), critical, warnings, info_count)
        else:
            cell.text = self._build_fallback_summary(len(insights), critical, warnings, info_count)

        for para in cell.paragraphs:
            for run in para.runs:
                run.font.size = Pt(11)
                run.font.name = "Calibri"

        doc.add_paragraph("")

    @staticmethod
    def _build_fallback_summary(total: int, critical: int, warnings: int, info: int) -> str:
        parts = [
            f"This analysis identified {total} areas of concern across the dataset.",
            "",
        ]
        if critical > 0:
            parts.append(f"\u2022 {critical} item(s) need immediate attention")
        if warnings > 0:
            parts.append(f"\u2022 {warnings} item(s) worth investigating further")
        if info > 0:
            parts.append(f"\u2022 {info} informational observation(s) for awareness")
        return "\n".join(parts)

    def _add_data_overview(
        self,
        doc: Document,
        insights: List[Dict[str, Any]],
        profiling: Optional[Dict[str, Any]],
    ) -> None:
        heading = doc.add_heading("Data Overview", level=1)
        heading.runs[0].font.color.rgb = NAVY

        overview_items = [
            ("Total Findings", str(len(insights))),
            ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ]

        if profiling:
            columns_info = profiling.get("columns", {})
            health = profiling.get("quality_metrics", {})
            overview_items.extend([
                ("Total Columns Profiled", str(len(columns_info))),
                ("Health Score", f"{health.get('health_score', 0):.1f}%"),
            ])

        resources = set()
        severities: Dict[str, int] = {}
        for insight in insights:
            resources.add(insight.get("resource", "N/A"))
            sev = str(insight.get("severity", "UNKNOWN")).upper()
            severities[sev] = severities.get(sev, 0) + 1

        overview_items.append(("Affected Resources", str(len(resources))))
        for sev, count in sorted(severities.items()):
            overview_items.append((f"  {sev} Findings", str(count)))

        table = doc.add_table(rows=len(overview_items) + 1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for j, header_text in enumerate(["Metric", "Value"]):
            cell = table.cell(0, j)
            cell.text = header_text
            _set_cell_bg(cell, '1F4E78')
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = WHITE

        for i, (metric, value) in enumerate(overview_items, start=1):
            table.cell(i, 0).text = metric
            table.cell(i, 1).text = value

        doc.add_paragraph("")

    def _add_key_findings(self, doc: Document, insights: List[Dict[str, Any]]) -> None:
        heading = doc.add_heading("Key Findings", level=1)
        heading.runs[0].font.color.rgb = NAVY

        severity_order = {"CRITICAL": 0, "WARNING": 1, "INFO": 2}
        sorted_insights = sorted(
            insights,
            key=lambda x: severity_order.get(str(x.get("severity", "INFO")).upper(), 3),
        )

        if not sorted_insights:
            doc.add_paragraph("No findings to report.")
            return

        sev_display = {
            "CRITICAL": "Needs Immediate Attention",
            "WARNING": "Worth Investigating",
            "INFO": "For Your Information",
        }

        display = sorted_insights[:25]
        table = doc.add_table(rows=len(display) + 1, cols=4)

        headers = ["#", "Priority", "Description", "Confidence"]
        for j, h in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = h
            _set_cell_bg(cell, '1F4E78')
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = WHITE

        for i, finding in enumerate(display, start=1):
            sev = str(finding.get("severity", "INFO")).upper()
            table.cell(i, 0).text = str(i)

            sev_cell = table.cell(i, 1)
            sev_cell.text = sev_display.get(sev, sev)
            color_map = {"CRITICAL": RED, "WARNING": AMBER, "INFO": GREEN}
            for run in sev_cell.paragraphs[0].runs:
                run.font.color.rgb = color_map.get(sev, GRAY)
                run.bold = True

            # FULL description — NEVER truncated
            desc = finding.get("description", finding.get("message", ""))
            table.cell(i, 2).text = str(desc)

            conf = finding.get("confidence", finding.get("priority_score", 0))
            table.cell(i, 3).text = f"{float(conf):.0%}" if conf else "N/A"

        if len(sorted_insights) > 25:
            doc.add_paragraph(
                f"... and {len(sorted_insights) - 25} additional findings. "
                "See full export for complete details."
            )

        doc.add_paragraph("")

    def _add_root_cause_table(self, doc: Document, insights: List[Dict[str, Any]]) -> None:
        heading = doc.add_heading("Root Cause Analysis", level=1)
        heading.runs[0].font.color.rgb = NAVY

        sev_display = {
            "CRITICAL": "Needs Attention",
            "WARNING": "Worth Investigating",
            "INFO": "For Your Information",
        }

        root_causes = []
        for insight in insights:
            remediation = insight.get("remediation", insight.get("action", ""))
            if remediation:
                root_causes.append({
                    "resource": insight.get("resource", "N/A"),
                    "finding": str(insight.get("description", insight.get("message", ""))),
                    "root_cause": str(remediation),
                    "severity": str(insight.get("severity", "INFO")).upper(),
                })

        if not root_causes:
            doc.add_paragraph("No root cause data available for current findings.")
            return

        display = root_causes[:20]
        table = doc.add_table(rows=len(display) + 1, cols=4)

        for j, h in enumerate(["Resource", "Finding", "Recommended Action", "Priority"]):
            cell = table.cell(0, j)
            cell.text = h
            _set_cell_bg(cell, '1F4E78')
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = WHITE

        for i, rc in enumerate(display, start=1):
            table.cell(i, 0).text = rc["resource"]
            table.cell(i, 1).text = rc["finding"]
            table.cell(i, 2).text = rc["root_cause"]
            sev_cell = table.cell(i, 3)
            sev_cell.text = sev_display.get(rc["severity"], rc["severity"])
            color_map = {"CRITICAL": RED, "WARNING": AMBER, "INFO": GREEN}
            for run in sev_cell.paragraphs[0].runs:
                run.font.color.rgb = color_map.get(rc["severity"], GRAY)
                run.bold = True

        doc.add_paragraph("")

    def _add_recommendations(self, doc: Document, insights: List[Dict[str, Any]]) -> None:
        heading = doc.add_heading("Recommendations", level=1)
        heading.runs[0].font.color.rgb = NAVY

        immediate: List[str] = []
        short_term: List[str] = []
        strategic: List[str] = []

        for insight in insights:
            sev = str(insight.get("severity", "INFO")).upper()
            action = insight.get("remediation", insight.get("action", ""))
            if not action:
                desc = insight.get("description", insight.get("message", ""))
                if desc:
                    action = f"Review and address: {desc}"
                else:
                    continue

            action_str = str(action)
            # Strip "[variable_name]" prefix pattern — show full sentence only
            if action_str.startswith("[") and "]" in action_str:
                action_str = action_str[action_str.index("]") + 1:].strip()
                if not action_str:
                    continue

            if sev == "CRITICAL":
                immediate.append(action_str)
            elif sev == "WARNING":
                short_term.append(action_str)
            else:
                strategic.append(action_str)

        self._add_recommendation_tier(
            doc, "Tier 1 \u2014 Immediate Action Required", immediate, RED
        )
        self._add_recommendation_tier(
            doc, "Tier 2 \u2014 Investigate This Week", short_term, AMBER
        )
        self._add_recommendation_tier(
            doc, "Tier 3 \u2014 Strategic Improvement", strategic, BLUE
        )

        doc.add_paragraph("")

    @staticmethod
    def _add_recommendation_tier(
        doc: Document,
        title: str,
        items: List[str],
        color: RGBColor,
    ) -> None:
        heading = doc.add_heading(title, level=2)
        for run in heading.runs:
            run.font.color.rgb = color

        if not items:
            p = doc.add_paragraph("No actions required at this level.")
            p.runs[0].font.color.rgb = GRAY
            p.runs[0].italic = True
            return

        for item in items[:10]:
            para = doc.add_paragraph(style="List Bullet")
            run = para.add_run(item)
            run.font.size = Pt(10)

        if len(items) > 10:
            p = doc.add_paragraph(f"  ... and {len(items) - 10} more")
            p.runs[0].font.color.rgb = GRAY

    def _add_predictive_signals(
        self,
        doc: Document,
        ml_findings: List[Dict[str, Any]],
    ) -> None:
        heading = doc.add_heading("Predictive Signals", level=1)
        heading.runs[0].font.color.rgb = NAVY

        if not ml_findings:
            doc.add_paragraph("No predictive signals generated for this run.")
            return

        type_labels = {
            "PREDICTION": "Prediction",
            "FORECAST": "Forecast",
            "RISK": "Risk Assessment",
            "ANOMALY": "Anomaly Detection",
        }
        sev_display = {
            "CRITICAL": "Needs Attention",
            "WARNING": "Worth Investigating",
            "INFO": "For Your Information",
        }

        for finding in ml_findings[:10]:
            finding_type = finding.get("type", "PREDICTION")
            description = finding.get("description", "")
            confidence = finding.get("confidence", 0)
            severity = str(finding.get("severity", "INFO")).upper()

            para = doc.add_paragraph()
            run = para.add_run(f"{type_labels.get(finding_type, finding_type)}: ")
            run.bold = True
            color_map = {"CRITICAL": RED, "WARNING": AMBER, "INFO": GREEN}
            run.font.color.rgb = color_map.get(severity, GRAY)

            run = para.add_run(str(description))
            run.font.size = Pt(10)

            conf_val = float(confidence) if confidence else 0
            if conf_val >= 0.8:
                conf_label = "High confidence"
            elif conf_val >= 0.5:
                conf_label = "Moderate confidence"
            else:
                conf_label = "Low confidence"

            detail = doc.add_paragraph()
            sev_label = sev_display.get(severity, severity)
            run = detail.add_run(f"   {conf_label} | {sev_label}")
            run.font.size = Pt(9)
            run.font.color.rgb = GRAY

        doc.add_paragraph("")

    @staticmethod
    def _add_footer(doc: Document) -> None:
        date_str = datetime.now().strftime("%B %d, %Y")
        footer_text = f"Confidential | Offline Industrial Intelligence | {date_str}"

        for section in doc.sections:
            footer = section.footer
            footer.is_linked_to_previous = False
            para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            para.text = footer_text
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.size = Pt(8)
                run.font.color.rgb = GRAY
